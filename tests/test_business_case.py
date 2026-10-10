"""Tests for scoring/business_case.py. All data is fictional."""

from __future__ import annotations

import os
import shutil
import subprocess
import sys
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from scoring import business_case as bc
from scoring.errors import InputError

TEMPLATE_INPUTS = Path(__file__).resolve().parent.parent / "templates" / "business_case_inputs.yaml"


@pytest.fixture(scope="module")
def case() -> bc.CaseInputs:
    return bc.load_inputs(TEMPLATE_INPUTS)


@pytest.fixture()
def raw() -> dict:
    return yaml.safe_load(TEMPLATE_INPUTS.read_text(encoding="utf-8"))


def write(tmp_path: Path, raw: dict) -> Path:
    path = tmp_path / "inputs.yaml"
    path.write_text(yaml.safe_dump(raw, sort_keys=False), encoding="utf-8")
    return path


# ---------------------------------------------------------------------------
# Python model, hand-calculated on the template (UC-01, fictional)
# ---------------------------------------------------------------------------


def test_base_case_hand_calculated(case):
    r = bc.compute(case.values("base"))
    assert r["hours_saved_per_month"] == pytest.approx(200)        # 3,000 x (6 - 2) / 60
    assert r["value_of_hours_per_month"] == pytest.approx(7_000)   # 200 x 50% x 70
    assert r["value_of_errors_per_month"] == pytest.approx(3_000)  # 3,000 x 2% x 50
    assert r["annual_benefit_steady"] == pytest.approx(96_000)     # 12 x 10,000 x 80%
    assert r["one_off_cost"] == pytest.approx(51_000)              # 60 x 600 + 15,000
    assert r["annual_run_cost"] == pytest.approx(30_000)           # 18,000 + 12,000


def test_base_case_payback_npv_roi(case):
    r = bc.compute(case.values("base"))
    # Net: -51,000 at month 0; 5,000 - 2,500 = 2,500 in months 1-3; 8,000 - 2,500 = 5,500 after.
    # Cumulative after month 3: -43,500; after month 10: -5,000; after month 11: +500.
    assert r["payback_month"] == 11
    # 3 x 5,000 + 33 x 8,000 = 279,000 benefit; 51,000 + 36 x 2,500 = 141,000 cost.
    assert r["total_benefit_3y"] == pytest.approx(279_000)
    assert r["total_cost_3y"] == pytest.approx(141_000)
    assert r["roi_3y"] == pytest.approx(138_000 / 141_000)
    expected_npv = -51_000 + sum((2_500 if m <= 3 else 5_500) / 1.08 ** (m / 12) for m in range(1, 37))
    assert r["npv_3y"] == pytest.approx(expected_npv)
    assert r["npv_3y"] == pytest.approx(116_352.91, abs=0.01)  # same value as Excel recalculation


def test_scenarios_change_only_the_three_uncertain_inputs(case):
    base, pess = case.values("base"), case.values("pessimistic")
    changed = {k for k in base if base[k] != pess[k]}
    assert changed == {"minutes_after", "share_hours_redeployed", "adoption_steady"}


def test_scenario_results(case):
    results = bc.compute_all(case)
    assert results["pessimistic"]["payback_month"] is None
    assert results["optimistic"]["payback_month"] == 7
    assert (results["pessimistic"]["npv_3y"] < results["base"]["npv_3y"]
            < results["optimistic"]["npv_3y"])


def test_hours_are_not_cash_unless_redeployed(case):
    values = dict(case.values("base"), share_hours_redeployed=0.0)
    r = bc.compute(values)
    assert r["value_of_hours_per_month"] == 0
    assert r["hours_not_counted_value_per_year"] == pytest.approx(200 * 70 * 12 * 0.8)


def test_benefit_follows_adoption_curve(case):
    months = bc.compute(case.values("base"))["months"]
    assert months[0]["benefit"] == 0 and months[0]["cost"] == 51_000
    assert months[3]["adoption"] == 0.5 and months[4]["adoption"] == 0.8
    assert len(months) == bc.HORIZON_MONTHS + 1


# ---------------------------------------------------------------------------
# Input validation
# ---------------------------------------------------------------------------


def test_missing_input_is_reported(tmp_path, raw):
    del raw["inputs"]["discount_rate"]
    with pytest.raises(InputError, match="missing input.*discount_rate"):
        bc.load_inputs(write(tmp_path, raw))


@pytest.mark.parametrize("key, field, value, message", [
    ("adoption_steady", "value", 1.5, "share and must be between 0 and 1"),
    ("build_days", "value", -5, "must not be negative"),
    ("build_days", "value", "sixty", "must be a number"),
    ("volume_per_month", "confidence", "certain", "confidence must be low, medium or high"),
    ("volume_per_month", "source", "", "needs a 'source'"),
])
def test_invalid_inputs_are_rejected(tmp_path, raw, key, field, value, message):
    raw["inputs"][key][field] = value
    with pytest.raises(InputError, match=message):
        bc.load_inputs(write(tmp_path, raw))


def test_exactly_three_scenario_inputs_are_required(tmp_path, raw):
    del raw["scenarios"]["minutes_after"]
    with pytest.raises(InputError, match="exactly three scenario inputs"):
        bc.load_inputs(write(tmp_path, raw))


def test_cli_returns_error_code_on_invalid_input(tmp_path, raw, capsys):
    raw["inputs"]["ramp_months"]["value"] = 2.5
    assert bc.main([str(write(tmp_path, raw))]) == 2
    assert "ramp_months must be a whole number" in capsys.readouterr().err


# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------


@pytest.fixture(scope="module")
def workbook_path(case, tmp_path_factory) -> Path:
    return bc.build_workbook(case, tmp_path_factory.mktemp("bc") / "business_case.xlsx")


def test_workbook_has_expected_sheets(workbook_path):
    wb = load_workbook(workbook_path)
    assert wb.sheetnames == ["Read me", "Inputs", "Scenarios", "Calculation", "Cash flow", "Assumptions log"]
    assert wb.calculation.fullCalcOnLoad
    assert not any(ws.protection.sheet for ws in wb.worksheets)


def test_every_input_has_a_defined_name_pointing_to_a_number(workbook_path):
    wb = load_workbook(workbook_path)
    for key in bc.INPUT_SPEC:
        assert key in wb.defined_names, key
        sheet, ref = list(wb.defined_names[key].destinations)[0]
        assert isinstance(wb[sheet][ref.replace("$", "")].value, (int, float)), key


def test_calculation_and_cash_flow_cells_are_formulas_not_numbers(workbook_path):
    wb = load_workbook(workbook_path)
    calc = wb["Calculation"]
    calc_cells = [calc.cell(row=r, column=c).value for r in range(6, calc.max_row + 1) for c in (3, 4, 5)]
    calc_values = [v for v in calc_cells if v is not None]
    assert len(calc_values) == 3 * (len(bc.INPUT_SPEC) + len(bc.CALC_ROWS))
    assert all(isinstance(v, str) and v.startswith("=") for v in calc_values)
    flow = wb["Cash flow"]
    flow_values = [flow.cell(row=r, column=c).value for r in range(5, 42) for c in range(2, 23)]
    assert all(isinstance(v, str) and v.startswith("=") for v in flow_values)


def test_scenario_sheet_holds_three_inputs_and_links_base(workbook_path, case):
    ws = load_workbook(workbook_path)["Scenarios"]
    ids = [ws.cell(row=r, column=1).value for r in range(5, 8)]
    assert ids == list(case.scenarios)
    assert all(str(ws.cell(row=r, column=4).value).startswith("=Inputs!") for r in range(5, 8))


def test_assumptions_log_covers_every_input_with_confidence_dropdown(workbook_path):
    ws = load_workbook(workbook_path)["Assumptions log"]
    ids = [ws.cell(row=r, column=1).value for r in range(5, 5 + len(bc.INPUT_SPEC))]
    assert ids == list(bc.INPUT_SPEC)
    assert all(ws.cell(row=r, column=6).value in bc.CONFIDENCE_LEVELS for r in range(5, 5 + len(bc.INPUT_SPEC)))
    assert any('"low,medium,high"' in dv.formula1 for dv in ws.data_validations.dataValidation)


def test_workbook_states_the_counting_rules_and_fiction(workbook_path):
    ws = load_workbook(workbook_path)["Read me"]
    text = " ".join(str(c.value) for c in ws["A"] if c.value)
    assert "FICTIONAL" in text
    assert "Hours saved are not cash" in text
    assert "A benefit counts only once" in text


def test_committed_template_workbook_matches_generator(case, tmp_path):
    """templates/business_case.xlsx must be regenerated after changing the inputs or the model."""
    committed = load_workbook(TEMPLATE_INPUTS.parent / "business_case.xlsx")
    fresh = load_workbook(bc.build_workbook(case, tmp_path / "fresh.xlsx"))
    for name in fresh.sheetnames:
        a, b = committed[name], fresh[name]
        assert [[c.value for c in row] for row in a.iter_rows()] == \
               [[c.value for c in row] for row in b.iter_rows()], name


@pytest.mark.skipif(os.environ.get("RUN_EXCEL_TESTS") != "1" or sys.platform != "win32" or not shutil.which("powershell"),
                    reason="Set RUN_EXCEL_TESTS=1 on Windows with Microsoft Excel installed.")
def test_excel_recalculation_matches_python(case, workbook_path):
    """Optional: recalculate the workbook in Excel (read-only) and compare with the Python model."""
    script = (
        "$xl = New-Object -ComObject Excel.Application; $xl.Visible=$false; $xl.DisplayAlerts=$false;"
        f"$wb = $xl.Workbooks.Open('{workbook_path}', 0, $true); $xl.CalculateFull();"
        "$ws = $wb.Worksheets.Item('Calculation');"
        "for ($r = 6; $r -le $ws.UsedRange.Rows.Count + 5; $r++) {"
        "  $l = $ws.Cells.Item($r,1).Text; if ($l) { Write-Output ($l + '|' + $ws.Cells.Item($r,4).Value2) } };"
        "$wb.Close($false); $xl.Quit()"
    )
    out = subprocess.run(["powershell", "-NoProfile", "-Command", script],
                         capture_output=True, text=True, timeout=120, check=True).stdout
    excel = dict(line.split("|", 1) for line in out.splitlines() if "|" in line)
    base = bc.compute(case.values("base"))
    assert float(excel["3-year NPV"]) == pytest.approx(base["npv_3y"])
    assert float(excel["Annual benefit at steady adoption"]) == pytest.approx(base["annual_benefit_steady"])
    assert float(excel["Payback (months)"]) == base["payback_month"]


# ---------------------------------------------------------------------------
# Markdown summary and CLI
# ---------------------------------------------------------------------------


def test_markdown_summary(case):
    md = bc.render_markdown(case)
    assert "FICTIONAL EXAMPLE" in md
    assert "| Payback (months) | not within 36 months | 11 | 7 |" in md
    assert "| Annual benefit at steady adoption | €40,500 | €96,000 | €159,975 |" in md
    assert "Hours saved are not cash" in md
    for key in bc.INPUT_SPEC:
        assert case.inputs[key]["label"] in md


def test_cli_writes_xlsx_and_md(tmp_path, capsys):
    xlsx, md = tmp_path / "bc.xlsx", tmp_path / "bc.md"
    assert bc.main([str(TEMPLATE_INPUTS), "--xlsx", str(xlsx), "--md", str(md)]) == 0
    assert xlsx.is_file() and md.is_file()
    assert "base: annual benefit €96,000, payback 11 months" in capsys.readouterr().out
