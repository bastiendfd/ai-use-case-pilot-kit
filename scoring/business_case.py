"""Business case model: Python reference calculation, Excel workbook and summary.

Usage (from the repository root):

    python -m scoring.business_case templates/business_case_inputs.yaml \
        --xlsx templates/business_case.xlsx --md path/to/business_case.md

The same model is implemented twice, on purpose:
- in Python (``compute``), for tests and for the Markdown summary;
- as live Excel formulas (``build_workbook``), so that a user can change any
  input in the workbook and see the result. No result is hard-coded.

Model (36 months, three scenarios that change only the three most uncertain
inputs):
- hours saved per month at full adoption = volume x (minutes before - after) / 60
- cash value of hours = hours saved x share redeployed x hourly cost
  (hours are cash only if they are redeployed to valued work or avoid a cost)
- value of errors avoided = volume x (error rate before - after) x cost per error
- monthly benefit = (cash value of hours + value of errors avoided) x adoption
- adoption = initial rate during the ramp-up, steady rate afterwards
- one-off cost at month 0 = build days x day rate + other one-off costs
- monthly run cost = (licence + run cost per year) / 12
- payback = first month in which cumulative net cash flow is >= 0
- NPV = sum of monthly net cash flows discounted at (1 + rate)^(month / 12)
- simple ROI = (total benefit - total cost) / total cost over 36 months
"""

from __future__ import annotations

import argparse
import sys
from dataclasses import dataclass
from pathlib import Path

import yaml

from scoring.errors import InputError

HORIZON_MONTHS = 36
SCENARIOS = ("pessimistic", "base", "optimistic")
CONFIDENCE_LEVELS = ("low", "medium", "high")

# Input id -> (group, number format). The order is the order in the workbook.
INPUT_SPEC = {
    "volume_per_month": ("Process", "#,##0"),
    "minutes_before": ("Process", "0.0"),
    "minutes_after": ("Process", "0.0"),
    "hourly_cost_eur": ("Process", "€#,##0"),
    "share_hours_redeployed": ("Process", "0%"),
    "error_rate_before": ("Process", "0.0%"),
    "error_rate_after": ("Process", "0.0%"),
    "cost_per_error_eur": ("Process", "€#,##0"),
    "licence_cost_per_year_eur": ("Costs", "€#,##0"),
    "build_days": ("Costs", "#,##0"),
    "day_rate_eur": ("Costs", "€#,##0"),
    "other_one_off_eur": ("Costs", "€#,##0"),
    "run_cost_per_year_eur": ("Costs", "€#,##0"),
    "adoption_initial": ("Adoption and finance", "0%"),
    "ramp_months": ("Adoption and finance", "0"),
    "adoption_steady": ("Adoption and finance", "0%"),
    "discount_rate": ("Adoption and finance", "0.0%"),
}
SHARE_INPUTS = ("share_hours_redeployed", "error_rate_before", "error_rate_after",
                "adoption_initial", "adoption_steady", "discount_rate")

FICTIONAL_NOTE = ("FICTIONAL EXAMPLE. All names and figures are invented for illustration "
                  "(\"Northwind Aero Services\").")

RULES = [
    "A benefit counts only once. If the same saved hours or avoided errors appear in another "
    "use case, count them in one case only.",
    "Hours saved are not cash. They count as cash only if they are redeployed to valued work "
    "or a cost is avoided (overtime, temporary staff, a planned hire). The share is an explicit input; "
    "the rest is shown as 'not counted'.",
    "Benefits ramp up with adoption. Nothing starts at 100% on day one.",
    "Run costs are part of the case: licences, hosting, monitoring, support, retraining.",
    "Scenarios change only the three most uncertain inputs. All other inputs stay at their base value.",
    "Every input has a source or reasoning and a confidence level in the assumptions log.",
    "A negative scenario is a decision input, not a footnote. If the 3-year NPV is negative in any "
    "scenario, the case can only be approved with conditions: a review gate 3 months after go-live, "
    "and stop or rework if adoption or time saved fall to their pessimistic values.",
]

# Downside rule (see RULES): review gate and the inputs checked at the gate.
REVIEW_GATE_MONTHS = 3
# "Time saved" is driven by minutes_after; "adoption" by adoption_steady.
GATE_DRIVERS_DEFAULT = ("minutes_after", "adoption_steady")


# ---------------------------------------------------------------------------
# Loading
# ---------------------------------------------------------------------------


@dataclass
class CaseInputs:
    use_case_id: str
    title: str
    fictional: bool
    inputs: dict[str, dict]        # id -> {label, value, unit, source, confidence}
    scenarios: dict[str, dict]     # id -> {pessimistic, optimistic, why}
    # Optional: scenario inputs checked at the review gate (default: GATE_DRIVERS_DEFAULT).
    gate_drivers: tuple[str, ...] | None = None

    def values(self, scenario: str = "base") -> dict[str, float]:
        """Input values for one scenario."""
        vals = {k: float(v["value"]) for k, v in self.inputs.items()}
        if scenario != "base":
            for k, s in self.scenarios.items():
                vals[k] = float(s[scenario])
        return vals


def load_inputs(path: Path) -> CaseInputs:
    """Read and validate a business case inputs YAML file."""
    path = Path(path)
    if not path.is_file():
        raise InputError([f"Business case inputs not found: {path}"])
    raw = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    errors = []
    inputs = raw.get("inputs") or {}
    scenarios = raw.get("scenarios") or {}

    missing = [k for k in INPUT_SPEC if k not in inputs]
    unknown = [k for k in inputs if k not in INPUT_SPEC]
    if missing:
        errors.append(f"{path}: missing input(s): {', '.join(missing)}.")
    if unknown:
        errors.append(f"{path}: unknown input(s): {', '.join(unknown)}.")
    for key, item in inputs.items():
        if key not in INPUT_SPEC:
            continue
        for field in ("label", "value", "unit", "source", "confidence"):
            if item.get(field) in (None, ""):
                errors.append(f"{path}: input '{key}' needs a '{field}'.")
        value = item.get("value")
        if not isinstance(value, (int, float)) or isinstance(value, bool):
            errors.append(f"{path}: input '{key}' must be a number, got '{value}'.")
        elif value < 0:
            errors.append(f"{path}: input '{key}' must not be negative.")
        elif key in SHARE_INPUTS and value > 1:
            errors.append(f"{path}: input '{key}' is a share and must be between 0 and 1, got {value}.")
        if str(item.get("confidence", "")).lower() not in CONFIDENCE_LEVELS:
            errors.append(f"{path}: input '{key}' confidence must be low, medium or high.")

    if len(scenarios) != 3:
        errors.append(f"{path}: exactly three scenario inputs are required, got {len(scenarios)}.")
    for key, item in scenarios.items():
        if key not in INPUT_SPEC:
            errors.append(f"{path}: scenario input '{key}' is not a known input.")
            continue
        for field in ("pessimistic", "optimistic", "why"):
            if item.get(field) in (None, ""):
                errors.append(f"{path}: scenario '{key}' needs '{field}'.")
        for field in ("pessimistic", "optimistic"):
            v = item.get(field)
            if v is not None and (not isinstance(v, (int, float)) or v < 0
                                  or (key in SHARE_INPUTS and v > 1)):
                errors.append(f"{path}: scenario '{key}' {field} value '{v}' is out of range.")

    gate_drivers = raw.get("review_gate_drivers")
    if gate_drivers is not None:
        if not isinstance(gate_drivers, list) or not gate_drivers:
            errors.append(f"{path}: review_gate_drivers must be a non-empty list of scenario inputs.")
        else:
            bad = [str(k) for k in gate_drivers if k not in scenarios]
            if bad:
                errors.append(f"{path}: review_gate_drivers {', '.join(bad)} must be scenario inputs.")

    if not errors:
        vals = {k: float(v["value"]) for k, v in inputs.items()}
        if vals["minutes_after"] > vals["minutes_before"] and "minutes_after" not in scenarios:
            errors.append(f"{path}: minutes_after is greater than minutes_before.")
        if vals["ramp_months"] != int(vals["ramp_months"]) or not 0 <= vals["ramp_months"] <= HORIZON_MONTHS:
            errors.append(f"{path}: ramp_months must be a whole number from 0 to {HORIZON_MONTHS}.")
    if errors:
        raise InputError(errors)
    for item in inputs.values():
        item["confidence"] = str(item["confidence"]).lower()
    return CaseInputs(
        use_case_id=str(raw.get("use_case_id", "")),
        title=str(raw.get("title", "")),
        fictional=bool(raw.get("fictional", False)),
        inputs={k: inputs[k] for k in INPUT_SPEC},
        scenarios=scenarios,
        gate_drivers=tuple(gate_drivers) if gate_drivers else None,
    )


# ---------------------------------------------------------------------------
# Python reference model
# ---------------------------------------------------------------------------


def compute(v: dict[str, float]) -> dict:
    """Compute the business case for one set of input values."""
    hours_saved = v["volume_per_month"] * (v["minutes_before"] - v["minutes_after"]) / 60
    hours_cash = hours_saved * v["share_hours_redeployed"]
    hours_value = hours_cash * v["hourly_cost_eur"]
    errors_avoided = v["volume_per_month"] * (v["error_rate_before"] - v["error_rate_after"])
    errors_value = errors_avoided * v["cost_per_error_eur"]
    gross_full = hours_value + errors_value
    one_off = v["build_days"] * v["day_rate_eur"] + v["other_one_off_eur"]
    annual_run = v["licence_cost_per_year_eur"] + v["run_cost_per_year_eur"]

    months = []
    cumulative = 0.0
    payback = None
    for m in range(HORIZON_MONTHS + 1):
        if m == 0:
            adoption, cost = 0.0, one_off
        else:
            adoption = v["adoption_initial"] if m <= v["ramp_months"] else v["adoption_steady"]
            cost = annual_run / 12
        benefit = gross_full * adoption
        net = benefit - cost
        previous = cumulative
        cumulative += net
        if payback is None and m > 0 and cumulative >= 0 > previous:
            payback = m
        months.append({
            "month": m, "adoption": adoption, "benefit": benefit, "cost": cost, "net": net,
            "cumulative": cumulative, "discounted": net / (1 + v["discount_rate"]) ** (m / 12),
        })

    total_benefit = sum(r["benefit"] for r in months)
    total_cost = sum(r["cost"] for r in months)
    return {
        "hours_saved_per_month": hours_saved,
        "hours_counted_as_cash_per_month": hours_cash,
        "value_of_hours_per_month": hours_value,
        "errors_avoided_per_month": errors_avoided,
        "value_of_errors_per_month": errors_value,
        "gross_monthly_benefit_full_adoption": gross_full,
        "annual_benefit_steady": 12 * gross_full * v["adoption_steady"],
        "one_off_cost": one_off,
        "annual_run_cost": annual_run,
        "payback_month": payback,
        "npv_3y": sum(r["discounted"] for r in months),
        "total_benefit_3y": total_benefit,
        "total_cost_3y": total_cost,
        "roi_3y": (total_benefit - total_cost) / total_cost if total_cost else None,
        "hours_not_counted_value_per_year": (hours_saved - hours_cash) * v["hourly_cost_eur"] * 12
        * v["adoption_steady"],
        "months": months,
    }


def compute_all(case: CaseInputs) -> dict[str, dict]:
    return {sc: compute(case.values(sc)) for sc in SCENARIOS}


def downside(case: CaseInputs, results: dict[str, dict] | None = None,
             drivers: tuple[str, ...] | None = None) -> dict:
    """Downside view used by the review-gate rule.

    Returns:
    - ``negative_scenarios``: scenarios with a negative 3-year NPV;
    - ``pessimistic_npv``: NPV with all three scenario inputs at pessimistic;
    - ``drivers``: one entry per scenario input with its base and pessimistic
      value, whether a higher value is worse, the NPV when only this input is
      at its pessimistic value, and whether it is a stop/rework trigger;
    - ``npv_at_triggers``: NPV when only the trigger inputs are at pessimistic.

    Trigger inputs: ``drivers`` if given, else ``review_gate_drivers`` from
    the inputs file, else GATE_DRIVERS_DEFAULT (time saved and adoption)
    among the scenario inputs.
    """
    results = results or compute_all(case)
    base_values = case.values("base")
    drivers = drivers if drivers is not None else case.gate_drivers
    triggers = tuple(drivers) if drivers is not None else tuple(
        k for k in GATE_DRIVERS_DEFAULT if k in case.scenarios)
    unknown = [k for k in triggers if k not in case.scenarios]
    if unknown:
        raise InputError([f"Review-gate driver(s) {', '.join(unknown)} must be scenario inputs "
                          f"({', '.join(case.scenarios)})."])
    rows = []
    for key, s in case.scenarios.items():
        pess = float(s["pessimistic"])
        rows.append({
            "input": key,
            "label": case.inputs[key]["label"],
            "unit": case.inputs[key]["unit"],
            "base": base_values[key],
            "pessimistic": pess,
            "higher_is_worse": pess > base_values[key],
            "npv_if_only_this": compute(dict(base_values, **{key: pess}))["npv_3y"],
            "trigger": key in triggers,
        })
        if key == "minutes_after":
            rows[-1]["time_saved_base"] = base_values["minutes_before"] - base_values[key]
            rows[-1]["time_saved_pessimistic"] = base_values["minutes_before"] - pess
    at_triggers = dict(base_values, **{k: float(case.scenarios[k]["pessimistic"]) for k in triggers})
    return {
        "negative_scenarios": [sc for sc in SCENARIOS if results[sc]["npv_3y"] < 0],
        "pessimistic_npv": results["pessimistic"]["npv_3y"],
        "drivers": rows,
        "trigger_inputs": triggers,
        "npv_at_triggers": compute(at_triggers)["npv_3y"],
    }


def minutes(value: float) -> str:
    """'1 minute', '1.5 minutes', '3 minutes'."""
    return f"{_fmt('minutes_after', value)} minute{'' if value == 1 else 's'}"


def trigger_text(row: dict) -> str:
    """Stop/rework threshold for one driver, e.g. 'at or below 60%'.

    For handling time, the threshold is stated as time saved, which is what
    the rule is about: 'time saved at or below 3 minutes per transaction
    (handling time at or above 3 minutes)'.
    """
    side = "at or above" if row["higher_is_worse"] else "at or below"
    if "time_saved_pessimistic" in row:
        return (f"time saved at or below {minutes(row['time_saved_pessimistic'])} per transaction "
                f"(handling time {side} {minutes(row['pessimistic'])})")
    unit = "" if row["input"] in SHARE_INPUTS else f" {row['unit']}"
    return f"{side} {_fmt(row['input'], row['pessimistic'])}{unit}"


# ---------------------------------------------------------------------------
# Excel workbook (live formulas)
# ---------------------------------------------------------------------------

# Calculation sheet rows: key -> (label, explanation, formula template, number format).
# The Calculation sheet starts with one row per input ("effective inputs"):
# the value from the Inputs sheet, or from the Scenarios sheet for the three
# scenario inputs. Formulas refer to those rows as {i_<input id>}.
# "{col}" is the scenario column (C, D or E) on the Calculation sheet;
# "{cf_<name>}" is a column of the matching block on the Cash flow sheet.
CALC_ROWS = [
    ("hours_saved", "Hours saved per month at full adoption", "volume x (minutes before - after) / 60",
     "={col}{i_volume_per_month}*({col}{i_minutes_before}-{col}{i_minutes_after})/60", "#,##0"),
    ("hours_cash", "  of which counted as cash", "hours saved x share redeployed",
     "={col}{r_hours_saved}*{col}{i_share_hours_redeployed}", "#,##0"),
    ("hours_value", "Cash value of redeployed hours per month", "hours counted as cash x hourly cost",
     "={col}{r_hours_cash}*{col}{i_hourly_cost_eur}", "€#,##0"),
    ("errors_avoided", "Errors avoided per month at full adoption", "volume x (error rate before - after)",
     "={col}{i_volume_per_month}*({col}{i_error_rate_before}-{col}{i_error_rate_after})", "#,##0"),
    ("errors_value", "Value of errors avoided per month", "errors avoided x cost per error",
     "={col}{r_errors_avoided}*{col}{i_cost_per_error_eur}", "€#,##0"),
    ("gross_full", "Monthly benefit at full adoption", "cash value of hours + value of errors avoided",
     "={col}{r_hours_value}+{col}{r_errors_value}", "€#,##0"),
    ("annual_benefit", "Annual benefit at steady adoption", "12 x monthly benefit x steady adoption",
     "=12*{col}{r_gross_full}*{col}{i_adoption_steady}", "€#,##0"),
    ("one_off", "One-off cost (month 0)", "build days x day rate + other one-off costs",
     "={col}{i_build_days}*{col}{i_day_rate_eur}+{col}{i_other_one_off_eur}", "€#,##0"),
    ("annual_run", "Annual run cost", "licence + run cost per year",
     "={col}{i_licence_cost_per_year_eur}+{col}{i_run_cost_per_year_eur}", "€#,##0"),
    ("payback", "Payback (months)", "first month with cumulative net cash flow >= 0",
     "=IF(COUNT('Cash flow'!{cf_payback}5:{cf_payback}41)=0,\"not within 36\",MIN('Cash flow'!{cf_payback}5:{cf_payback}41))", "0"),
    ("npv", "3-year NPV", "sum of monthly net cash flows discounted at (1 + rate)^(month / 12)",
     "=SUM('Cash flow'!{cf_discounted}5:{cf_discounted}41)", "€#,##0"),
    ("total_benefit", "Total benefit over 3 years", "sum of monthly benefits",
     "=SUM('Cash flow'!{cf_benefit}5:{cf_benefit}41)", "€#,##0"),
    ("total_cost", "Total cost over 3 years", "one-off + monthly run costs",
     "=SUM('Cash flow'!{cf_cost}5:{cf_cost}41)", "€#,##0"),
    ("roi", "Simple ROI over 3 years", "(total benefit - total cost) / total cost",
     "=({col}{r_total_benefit}-{col}{r_total_cost})/{col}{r_total_cost}", "0%"),
    ("not_counted", "Not counted: value of saved hours not redeployed, per year", "for information only; not in NPV",
     "=({col}{r_hours_saved}-{col}{r_hours_cash})*{col}{i_hourly_cost_eur}*12*{col}{i_adoption_steady}", "€#,##0"),
]
CASHFLOW_COLS = ("adoption", "benefit", "cost", "net", "cumulative", "discounted", "payback")


def _col(index: int) -> str:
    from openpyxl.utils import get_column_letter
    return get_column_letter(index)


def build_workbook(case: CaseInputs, path: Path) -> Path:
    """Write the business case workbook with live formulas."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation

    bold = Font(bold=True)
    title_font = Font(bold=True, size=14)
    input_font = Font(color="0000CC")
    input_fill = PatternFill("solid", fgColor="FFF2CC")
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    result_fill = PatternFill("solid", fgColor="E2EFDA")
    thin = Side(style="thin", color="BFBFBF")
    box = Border(top=thin, bottom=thin, left=thin, right=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    def header(ws, row, values):
        for i, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=i, value=value)
            cell.font, cell.fill, cell.border = bold, header_fill, box

    def banner(ws, title):
        ws["A1"] = title
        ws["A1"].font = title_font
        ws["A2"] = f"{case.use_case_id}: {case.title}"
        if case.fictional:
            ws["A3"] = FICTIONAL_NOTE
            ws["A3"].font = Font(italic=True, color="C00000")

    wb = Workbook()

    # Read me ------------------------------------------------------------------
    ws = wb.active
    ws.title = "Read me"
    banner(ws, "Business case")
    lines = [
        "",
        "How to use",
        "1. Enter inputs on the 'Inputs' sheet (blue text on yellow = input). Everything else is a formula.",
        "2. Enter the pessimistic and optimistic values of the three most uncertain inputs on 'Scenarios'.",
        "3. Record the source and confidence of every input on 'Assumptions log'.",
        "4. Read the results on 'Calculation'. 'Cash flow' shows the 36 months behind payback and NPV.",
        "",
        "Counting rules",
        *[f"- {r}" for r in RULES],
        "",
        "Nothing in this workbook is protected. Results are recalculated when the file is opened.",
    ]
    for i, text in enumerate(lines, start=5):
        ws.cell(row=i, column=1, value=text)
        if text in ("How to use", "Counting rules"):
            ws.cell(row=i, column=1).font = bold
    ws.column_dimensions["A"].width = 120

    # Inputs ---------------------------------------------------------------------
    ws = wb.create_sheet("Inputs")
    banner(ws, "Inputs")
    header(ws, 4, ["Group", "ID (defined name)", "Input", "Value", "Unit"])
    input_row = {}
    for r, (key, (group, fmt)) in enumerate(INPUT_SPEC.items(), start=5):
        item = case.inputs[key]
        ws.cell(row=r, column=1, value=group)
        ws.cell(row=r, column=2, value=key)
        ws.cell(row=r, column=3, value=item["label"])
        cell = ws.cell(row=r, column=4, value=item["value"])
        cell.number_format, cell.font, cell.fill, cell.border = fmt, input_font, input_fill, box
        ws.cell(row=r, column=5, value=item["unit"] + (" (base; see Scenarios)" if key in case.scenarios else ""))
        input_row[key] = r
        wb.defined_names[key] = DefinedName(key, attr_text=f"Inputs!$D${r}")
    for col, width in zip("ABCDE", (22, 28, 52, 14, 40)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A5"

    # Scenarios ------------------------------------------------------------------
    ws = wb.create_sheet("Scenarios")
    banner(ws, "Scenarios: the three most uncertain inputs")
    header(ws, 4, ["ID", "Input", "Pessimistic", "Base", "Optimistic", "Unit", "Why it is uncertain"])
    scen_row = {}
    for r, (key, item) in enumerate(case.scenarios.items(), start=5):
        fmt = INPUT_SPEC[key][1]
        ws.cell(row=r, column=1, value=key)
        ws.cell(row=r, column=2, value=case.inputs[key]["label"])
        for c, value in ((3, item["pessimistic"]), (5, item["optimistic"])):
            cell = ws.cell(row=r, column=c, value=value)
            cell.number_format, cell.font, cell.fill, cell.border = fmt, input_font, input_fill, box
        base = ws.cell(row=r, column=4, value=f"=Inputs!$D${input_row[key]}")
        base.number_format, base.border = fmt, box
        ws.cell(row=r, column=6, value=case.inputs[key]["unit"])
        ws.cell(row=r, column=7, value=item["why"]).alignment = wrap
        scen_row[key] = r
    ws.cell(row=9, column=1, value="Base values come from the Inputs sheet. Change them there.").font = Font(italic=True)
    for col, width in zip("ABCDEFG", (24, 46, 13, 13, 13, 22, 60)):
        ws.column_dimensions[col].width = width

    # Calculation ----------------------------------------------------------------
    ws = wb.create_sheet("Calculation")
    banner(ws, "Calculation")
    header(ws, 4, ["Metric", "How it is calculated", "Pessimistic", "Base", "Optimistic"])
    ws.cell(row=5, column=1, value="Effective inputs").font = bold
    input_calc_row = {key: r for r, key in enumerate(INPUT_SPEC, start=6)}
    first_result = 6 + len(INPUT_SPEC) + 1
    ws.cell(row=first_result, column=1, value="Results").font = bold
    calc_row = {key: r for r, (key, *_rest) in enumerate(CALC_ROWS, start=first_result + 1)}
    cf_first = {sc: 2 + i * len(CASHFLOW_COLS) for i, sc in enumerate(SCENARIOS)}
    refs = {f"r_{k}": v for k, v in calc_row.items()}
    refs.update({f"i_{k}": v for k, v in input_calc_row.items()})
    for key, r in input_calc_row.items():
        fmt = INPUT_SPEC[key][1]
        ws.cell(row=r, column=1, value=case.inputs[key]["label"])
        ws.cell(row=r, column=2, value="Scenarios sheet" if key in case.scenarios else "Inputs sheet (same in all scenarios)")
        for i in range(len(SCENARIOS)):
            value = f"=Scenarios!{_col(3 + i)}{scen_row[key]}" if key in case.scenarios else f"={key}"
            cell = ws.cell(row=r, column=3 + i, value=value)
            cell.number_format, cell.border = fmt, box
            if key in case.scenarios:
                cell.font = bold
    results_keys = {"annual_benefit", "one_off", "annual_run", "payback", "npv", "roi"}
    for key, label, how, template, fmt in CALC_ROWS:
        r = calc_row[key]
        ws.cell(row=r, column=1, value=label).font = bold if key in results_keys else Font()
        ws.cell(row=r, column=2, value=how)
        for i, sc in enumerate(SCENARIOS):
            col = _col(3 + i)
            first = cf_first[sc]
            cf = {f"cf_{name}": _col(first + j) for j, name in enumerate(CASHFLOW_COLS)}
            formula = template.format(col=col, **refs, **cf)
            cell = ws.cell(row=r, column=3 + i, value=formula)
            cell.number_format, cell.border = fmt, box
            if key in results_keys:
                cell.fill = result_fill
    ws.cell(row=max(calc_row.values()) + 2, column=1,
            value="Hours are counted as cash only for the redeployed share. See the counting rules on 'Read me'.").font = Font(italic=True)
    for col, width in zip("ABCDE", (56, 50, 15, 15, 15)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "C5"

    # Cash flow ------------------------------------------------------------------
    ws = wb.create_sheet("Cash flow")
    banner(ws, "Monthly cash flow (36 months)")
    ws.cell(row=4, column=1, value="Month").font = bold
    ws.cell(row=4, column=1).fill = header_fill
    labels = {"adoption": "Adoption", "benefit": "Benefit", "cost": "Cost", "net": "Net",
              "cumulative": "Cumulative", "discounted": "Discounted net", "payback": "Payback month"}
    formats = {"adoption": "0%", "benefit": "€#,##0", "cost": "€#,##0", "net": "€#,##0",
               "cumulative": "€#,##0", "discounted": "€#,##0", "payback": "0"}
    for sc in SCENARIOS:
        first = cf_first[sc]
        ws.cell(row=3, column=first, value=sc.capitalize()).font = bold
        ws.merge_cells(start_row=3, start_column=first, end_row=3, end_column=first + len(CASHFLOW_COLS) - 1)
        for j, name in enumerate(CASHFLOW_COLS):
            cell = ws.cell(row=4, column=first + j, value=labels[name])
            cell.font, cell.fill, cell.border = bold, header_fill, box
    for m in range(HORIZON_MONTHS + 1):
        r = 5 + m
        ws.cell(row=r, column=1, value=m)
        for i, sc in enumerate(SCENARIOS):
            first = cf_first[sc]
            c = {name: _col(first + j) for j, name in enumerate(CASHFLOW_COLS)}
            calc = _col(3 + i)
            f = {
                "adoption": f"=IF($A{r}=0,0,IF($A{r}<=Calculation!{calc}${input_calc_row['ramp_months']},"
                            f"Calculation!{calc}${input_calc_row['adoption_initial']},"
                            f"Calculation!{calc}${input_calc_row['adoption_steady']}))",
                "benefit": f"=Calculation!{calc}${calc_row['gross_full']}*{c['adoption']}{r}",
                "cost": f"=IF($A{r}=0,Calculation!{calc}${calc_row['one_off']},Calculation!{calc}${calc_row['annual_run']}/12)",
                "net": f"={c['benefit']}{r}-{c['cost']}{r}",
                "cumulative": f"={c['net']}{r}" if m == 0 else f"={c['cumulative']}{r - 1}+{c['net']}{r}",
                "discounted": f"={c['net']}{r}/(1+Calculation!{calc}${input_calc_row['discount_rate']})^($A{r}/12)",
                "payback": '=""' if m == 0 else
                f"=IF(AND({c['cumulative']}{r}>=0,{c['cumulative']}{r - 1}<0),$A{r},\"\")",
            }
            for name in CASHFLOW_COLS:
                ws.cell(row=r, column=first + CASHFLOW_COLS.index(name), value=f[name]).number_format = formats[name]
    ws.column_dimensions["A"].width = 8
    for idx in range(2, 2 + len(SCENARIOS) * len(CASHFLOW_COLS)):
        ws.column_dimensions[_col(idx)].width = 13
    ws.freeze_panes = "B5"

    # Assumptions log ------------------------------------------------------------
    ws = wb.create_sheet("Assumptions log")
    banner(ws, "Assumptions log")
    header(ws, 4, ["ID", "Input", "Value (base)", "Unit", "Source or reasoning", "Confidence", "Scenario range"])
    dv = DataValidation(type="list", formula1='"low,medium,high"', allow_blank=False)
    ws.add_data_validation(dv)
    for r, key in enumerate(INPUT_SPEC, start=5):
        item = case.inputs[key]
        ws.cell(row=r, column=1, value=key)
        ws.cell(row=r, column=2, value=item["label"])
        cell = ws.cell(row=r, column=3, value=f"={key}")
        cell.number_format = INPUT_SPEC[key][1]
        ws.cell(row=r, column=4, value=item["unit"])
        ws.cell(row=r, column=5, value=item["source"]).alignment = wrap
        conf = ws.cell(row=r, column=6, value=item["confidence"])
        conf.font, conf.fill = input_font, input_fill
        dv.add(conf)
        if key in case.scenarios:
            s = case.scenarios[key]
            ws.cell(row=r, column=7, value=f"{s['pessimistic']} / base / {s['optimistic']}")
    for col, width in zip("ABCDEFG", (26, 46, 14, 22, 70, 12, 22)):
        ws.column_dimensions[col].width = width

    wb.calculation.fullCalcOnLoad = True
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# Markdown summary
# ---------------------------------------------------------------------------


def _eur(x: float) -> str:
    sign = "-" if x < 0 else ""
    return f"{sign}€{abs(x):,.0f}"


def _fmt(key: str, value: float) -> str:
    if key in SHARE_INPUTS:
        return f"{value:.0%}" if value * 100 == int(value * 100) else f"{value:.1%}"
    return f"{value:,.0f}" if value == int(value) else f"{value:,.1f}"


def render_markdown(case: CaseInputs, results: dict[str, dict] | None = None) -> str:
    """Business case summary in the structure of templates/business_case.md."""
    results = results or compute_all(case)
    base = results["base"]

    def row(label, key, fmt=_eur):
        cells = []
        for sc in SCENARIOS:
            v = results[sc][key]
            cells.append("not within 36 months" if v is None else fmt(v))
        return f"| {label} | " + " | ".join(cells) + " |"

    lines = [f"# Business case: {case.use_case_id} {case.title}", ""]
    if case.fictional:
        lines += [f"> **{FICTIONAL_NOTE}**", ""]
    lines += [
        "Generated by `scoring/business_case.py` from `business_case_inputs.yaml` in this folder. "
        "The workbook `templates/business_case.xlsx` implements the same model with live formulas.",
        "",
        "## 1. Summary",
        "",
        "| Result | Pessimistic | Base | Optimistic |",
        "|---|---|---|---|",
        row("Annual benefit at steady adoption", "annual_benefit_steady"),
        row("One-off cost", "one_off_cost"),
        row("Annual run cost", "annual_run_cost"),
        row("Payback (months)", "payback_month", lambda v: f"{v:.0f}"),
        row("3-year NPV", "npv_3y"),
        row("Simple ROI over 3 years", "roi_3y", lambda v: f"{v:.0%}"),
        "",
        "## 2. Benefit logic (base scenario)",
        "",
        "| Benefit | Mechanism | Per month at full adoption |",
        "|---|---|---|",
        f"| Hours saved | {base['hours_saved_per_month']:,.0f} hours; only the redeployed share "
        f"({base['hours_counted_as_cash_per_month']:,.0f} hours) is counted as cash | {_eur(base['value_of_hours_per_month'])} |",
        f"| Errors avoided | {base['errors_avoided_per_month']:,.0f} errors avoided × cost per error | "
        f"{_eur(base['value_of_errors_per_month'])} |",
        f"| **Total** | Scaled by adoption: {_fmt('adoption_initial', case.values()['adoption_initial'])} during the "
        f"first {case.values()['ramp_months']:.0f} months, then {_fmt('adoption_steady', case.values()['adoption_steady'])} "
        f"| **{_eur(base['gross_monthly_benefit_full_adoption'])}** |",
        "",
        f"Not counted: saved hours that are not redeployed, worth about "
        f"{_eur(base['hours_not_counted_value_per_year'])} per year at steady adoption. They are shown for "
        "information only and are not in the NPV.",
        "",
        "## 3. Inputs and assumptions log",
        "",
        "| Input | Base value | Unit | Source or reasoning | Confidence |",
        "|---|---|---|---|---|",
    ]
    for key in INPUT_SPEC:
        item = case.inputs[key]
        lines.append(f"| {item['label']} | {_fmt(key, float(item['value']))} | {item['unit']} | "
                     f"{item['source']} | {item['confidence']} |")
    lines += [
        "",
        "## 4. Scenarios",
        "",
        "Only the three most uncertain inputs change.",
        "",
        "| Input | Pessimistic | Base | Optimistic | Why it is uncertain |",
        "|---|---|---|---|---|",
    ]
    for key, s in case.scenarios.items():
        lines.append(f"| {case.inputs[key]['label']} | {_fmt(key, s['pessimistic'])} | "
                     f"{_fmt(key, float(case.inputs[key]['value']))} | {_fmt(key, s['optimistic'])} | {s['why']} |")
    lines += _downside_markdown(case, results)
    lines += ["", "## 6. Counting rules applied", ""] + [f"- {r}" for r in RULES] + [""]
    return "\n".join(lines)


def _downside_markdown(case: CaseInputs, results: dict[str, dict]) -> list[str]:
    """Section 5 of the summary: the downside rule applied to this case."""
    d = downside(case, results)
    lines = ["", "## 5. Downside and review gate", ""]
    if not d["negative_scenarios"]:
        return lines + [
            f"The 3-year NPV is positive in every scenario (pessimistic: {_eur(d['pessimistic_npv'])}). "
            "The downside rule does not add a review gate; the standard value review applies.",
        ]
    lines += [
        f"**The 3-year NPV is negative in the {', '.join(d['negative_scenarios'])} scenario"
        f"{'s' if len(d['negative_scenarios']) > 1 else ''}. Pessimistic NPV: "
        f"{_eur(d['pessimistic_npv'])}.**",
        "",
    ]
    if "base" in d["negative_scenarios"]:
        lines += ["The base case itself is negative, so a review gate cannot make the case approvable: "
                  "rework the case (scope, cost, adoption) or reject it. For information, the effect of each "
                  "scenario input:", ""]
    else:
        lines += [
        f"Rule: the case can only be approved with conditions. A review gate takes place "
        f"{REVIEW_GATE_MONTHS} months after go-live (or after the pilot start). If a trigger input is "
        "at its pessimistic value or worse, the case is stopped or reworked; it continues only if the "
        "business case, re-run on the measured values, has a positive NPV.",
        "",
        ]
    lines += [
        "| Input | Base | Pessimistic | NPV if only this input is pessimistic | Stop/rework trigger |",
        "|---|---|---|---|---|",
    ]
    for row in d["drivers"]:
        lines.append(
            f"| {row['label']} | {_fmt(row['input'], row['base'])} | {_fmt(row['input'], row['pessimistic'])} "
            f"| {_eur(row['npv_if_only_this'])} | {trigger_text(row) if row['trigger'] else 'no (monitored)'} |")
    labels = "; ".join(r["label"] for r in d["drivers"] if r["trigger"])
    lines += [
        "",
        f"NPV if only the trigger inputs ({labels}) are at their pessimistic values: "
        f"{_eur(d['npv_at_triggers'])}. NPV with all three scenario inputs at pessimistic: "
        f"{_eur(d['pessimistic_npv'])}.",
    ]
    return lines


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="python -m scoring.business_case",
                                     description="Build a business case workbook and summary.")
    parser.add_argument("inputs", type=Path, help="Business case inputs YAML.")
    parser.add_argument("--xlsx", type=Path, default=None, help="Write the workbook here.")
    parser.add_argument("--md", type=Path, default=None, help="Write the Markdown summary here.")
    args = parser.parse_args(argv)
    try:
        case = load_inputs(args.inputs)
    except InputError as exc:
        print("Input error:", file=sys.stderr)
        for message in exc.errors:
            print(f"  - {message}", file=sys.stderr)
        return 2
    results = compute_all(case)
    for sc in SCENARIOS:
        r = results[sc]
        payback = "not within 36" if r["payback_month"] is None else f"{r['payback_month']} months"
        print(f"{sc:>11}: annual benefit {_eur(r['annual_benefit_steady'])}, payback {payback}, "
              f"NPV {_eur(r['npv_3y'])}, ROI {r['roi_3y']:.0%}")
    d = downside(case, results)
    if results["base"]["npv_3y"] < 0:
        print("Base case NPV is negative: do not approve as is; rework the case or reject it.")
    elif d["negative_scenarios"]:
        print(f"Negative NPV in: {', '.join(d['negative_scenarios'])}. Approve only with conditions: "
              f"review gate {REVIEW_GATE_MONTHS} months after go-live (see the summary, section 5).")
    if args.xlsx:
        build_workbook(case, args.xlsx)
        print(f"Wrote {args.xlsx}")
    if args.md:
        args.md.parent.mkdir(parents=True, exist_ok=True)
        args.md.write_text(render_markdown(case, results), encoding="utf-8")
        print(f"Wrote {args.md}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
